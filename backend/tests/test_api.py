import io
import pytest
from docx import Document
from app import create_app


@pytest.fixture(scope='module')
def client():
    app = create_app('testing')
    return app.test_client()


@pytest.fixture
def sample_docx():
    doc = Document()
    doc.add_paragraph('JANE DOE')
    doc.add_paragraph('jane.doe@email.com | +1 555-987-6543')
    doc.add_paragraph('SUMMARY')
    doc.add_paragraph('Data analyst with experience in SQL, Python, and Tableau.')
    doc.add_paragraph('SKILLS')
    doc.add_paragraph('SQL, Python, Pandas, Tableau, Power BI, Excel, Statistics, Data Visualization')
    doc.add_paragraph('EDUCATION')
    doc.add_paragraph('B.Sc. Statistics, City University (2015-2019)')
    doc.add_paragraph('EXPERIENCE')
    doc.add_paragraph('Data Analyst, AnalyticsCo (2020-Present) - Built dashboards and reports used by management.')
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


class TestHealthEndpoint:
    def test_health_check(self, client):
        resp = client.get('/api/health')
        assert resp.status_code == 200
        assert resp.get_json()['status'] == 'healthy'


class TestUploadEndpoint:
    def test_upload_valid_docx(self, client, sample_docx):
        resp = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_doe_resume.docx')},
            content_type='multipart/form-data'
        )
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['resume_id'] is not None
        assert data['file_type'] == 'docx'
        assert data['text_length'] > 0
        assert data['contact_info']['email'] == 'jane.doe@email.com'

    def test_upload_missing_file(self, client):
        resp = client.post('/api/upload', data={}, content_type='multipart/form-data')
        assert resp.status_code == 400
        assert 'error' in resp.get_json()

    def test_upload_unsupported_type(self, client):
        resp = client.post(
            '/api/upload',
            data={'file': (io.BytesIO(b'data'), 'resume.txt')},
            content_type='multipart/form-data'
        )
        assert resp.status_code == 400

    def test_upload_corrupted_pdf(self, client):
        resp = client.post(
            '/api/upload',
            data={'file': (io.BytesIO(b'not a real pdf'), 'resume.pdf')},
            content_type='multipart/form-data'
        )
        assert resp.status_code == 400


class TestAnalysisEndpoint:
    def test_analyze_valid_resume(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_resume.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        resp = client.post('/api/analyze', json={'resume_id': rid})
        assert resp.status_code == 200
        data = resp.get_json()
        assert 0 <= data['overall_score'] <= 100
        assert len(data['sections']) == 6

    def test_analyze_missing_resume_id(self, client):
        resp = client.post('/api/analyze', json={})
        assert resp.status_code == 400

    def test_analyze_nonexistent_resume(self, client):
        resp = client.post('/api/analyze', json={'resume_id': 99999})
        assert resp.status_code == 404


class TestATSEndpoint:
    def test_ats_valid_role(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_ats.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        resp = client.post('/api/ats', json={'resume_id': rid, 'target_role': 'Data Analyst'})
        assert resp.status_code == 200
        data = resp.get_json()
        assert 0 <= data['ats_score'] <= 100
        assert data['target_role'] == 'Data Analyst'
        assert 'matched_skills' in data
        assert 'missing_skills' in data

    def test_ats_invalid_role(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_ats2.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        resp = client.post('/api/ats', json={'resume_id': rid, 'target_role': 'Unicorn Role'})
        assert resp.status_code == 400

    def test_ats_missing_role(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_ats3.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        resp = client.post('/api/ats', json={'resume_id': rid})
        assert resp.status_code == 400


class TestFeedbackEndpoint:
    def test_feedback_valid(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_fb.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        resp = client.post('/api/feedback', json={'resume_id': rid, 'target_role': 'Data Analyst'})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data['total_suggestions'] > 0
        for item in data['feedback']:
            assert 'category' in item
            assert 'title' in item
            assert 'impact' in item


class TestDashboardEndpoint:
    def test_dashboard_returns_full_data(self, client, sample_docx):
        upload = client.post(
            '/api/upload',
            data={'file': (sample_docx, 'jane_dash.docx')},
            content_type='multipart/form-data'
        ).get_json()
        rid = upload['resume_id']

        analysis = client.post('/api/analyze', json={'resume_id': rid}).get_json()
        aid = analysis['analysis_id']

        client.post('/api/ats', json={'resume_id': rid, 'target_role': 'Data Analyst'})

        resp = client.get(f'/api/dashboard/{aid}')
        assert resp.status_code == 200
        data = resp.get_json()
        assert 'overall_score' in data
        assert 'ats_score' in data
        assert 'matched_skills' in data
        assert 'missing_skills' in data
        assert 'sections' in data
        assert 'feedback' in data

    def test_dashboard_nonexistent(self, client):
        resp = client.get('/api/dashboard/99999')
        assert resp.status_code == 404


class TestRolesEndpoint:
    def test_roles_list(self, client):
        resp = client.get('/api/roles')
        assert resp.status_code == 200
        roles = resp.get_json()['roles']
        assert 'Software Engineer' in roles
        assert 'Data Analyst' in roles
