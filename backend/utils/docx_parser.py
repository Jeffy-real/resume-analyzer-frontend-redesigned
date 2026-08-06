from docx import Document
from utils.text_cleaner import clean_text

def parse_docx(file_path: str) -> str:
    try:
        doc = Document(file_path)
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        full_text = '\n'.join(text_parts)
        cleaned_text = clean_text(full_text)

        if len(cleaned_text.strip()) < 50:
            raise ValueError("Extracted text is too short. The document may be empty or corrupted.")

        return cleaned_text

    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"DOCX parsing failed: {str(e)}")
